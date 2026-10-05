"""Private application records, verified career profiles and ephemeral CV processing."""
import io
import re
from datetime import date, timedelta, datetime, timezone
from typing import Literal
from urllib.parse import urlsplit
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field, field_validator
from database.supabase_client import supabase
from sessions import current_user, issue_session

router = APIRouter()

class Application(BaseModel):
    job_title: str = Field(min_length=1, max_length=250)
    company_name: str = Field(min_length=1, max_length=250)
    company_website: str = Field(default="", max_length=2000)
    location: str = Field(default="", max_length=250)
    source_url: str = Field(default="", max_length=2000)
    source: str = Field(default="manual", max_length=100)
    date_applied: date | None = None
    status: Literal["started", "applied", "interview", "offer", "rejected", "ghosted"] = "started"
    notes: str = Field(default="", max_length=10000)
    salary_min: float | None = Field(default=None, ge=0)
    salary_max: float | None = Field(default=None, ge=0)
    salary_currency: str = Field(default="", max_length=10)
    salary_period: str = Field(default="", max_length=30)
    contact_person: str = Field(default="", max_length=250)
    deadline: date | None = None
    reminder_date: date | None = None
    email_reminder: bool = False

    @field_validator("source_url", "company_website")
    @classmethod
    def safe_url(cls, value):
        if value and (urlsplit(value).scheme not in ("https", "http") or not urlsplit(value).netloc):
            raise ValueError("Use a complete http or https link.")
        return value

    def record(self):
        data = self.model_dump(mode="json")
        if self.status == "applied" and not self.date_applied:
            data["date_applied"] = date.today().isoformat()
        if self.deadline and not self.reminder_date:
            data["reminder_date"] = (self.deadline - timedelta(days=2)).isoformat()
        return data

@router.get("/account/me")
def me(user=Depends(current_user)):
    result = supabase.table("profiles").select("id,email,full_name,country").eq("id", user).execute()
    if not result.data: raise HTTPException(401, "Account not found.")
    return result.data[0]

@router.get("/account/applications")
def applications(user=Depends(current_user)):
    return {"applications": supabase.table("applications").select("*").eq("user_id", user).order("date_applied", desc=True, nullsfirst=False).order("created_at", desc=True).limit(500).execute().data}

@router.post("/account/applications")
def add_application(data: Application, user=Depends(current_user)):
    if data.source_url:
        existing = supabase.table("applications").select("*").eq("user_id", user).eq("source_url", data.source_url).limit(1).execute().data
        if existing: return existing[0]
    return supabase.table("applications").insert({**data.record(), "user_id": user}).execute().data[0]

@router.put("/account/applications/{application_id}")
def update_application(application_id: str, data: Application, user=Depends(current_user)):
    existing = supabase.table("applications").select("reminder_date,email_reminder").eq("id", application_id).eq("user_id", user).execute().data
    if not existing: raise HTTPException(404, "Application not found.")
    payload = {**data.record(), "updated_at": datetime.now(timezone.utc).isoformat()}
    if payload["reminder_date"] != existing[0].get("reminder_date") or payload["email_reminder"] != existing[0].get("email_reminder"):
        payload["email_reminder_sent_at"] = None
    result = supabase.table("applications").update(payload).eq("id", application_id).eq("user_id", user).execute()
    if not result.data: raise HTTPException(404, "Application not found.")
    return result.data[0]

@router.delete("/account/applications/{application_id}")
def delete_application(application_id: str, user=Depends(current_user)):
    result = supabase.table("applications").delete().eq("id", application_id).eq("user_id", user).execute()
    if not result.data: raise HTTPException(404, "Application not found.")
    return {"deleted": True}

@router.get("/account/applications/{application_id}/events")
def events(application_id: str, user=Depends(current_user)):
    owns = supabase.table("applications").select("id").eq("id", application_id).eq("user_id", user).execute()
    if not owns.data: raise HTTPException(404, "Application not found.")
    return {"events": supabase.table("application_events").select("*").eq("application_id", application_id).order("created_at", desc=True).execute().data}

class Profile(BaseModel):
    verified_text: str = Field(max_length=30000)

@router.get("/account/profile")
def get_profile(user=Depends(current_user)):
    rows = supabase.table("career_profiles").select("verified_text").eq("user_id", user).execute().data
    return rows[0] if rows else {"verified_text": ""}

@router.put("/account/profile")
def save_profile(data: Profile, user=Depends(current_user)):
    supabase.table("career_profiles").upsert({"user_id": user, "verified_text": data.verified_text}, on_conflict="user_id").execute()
    return {"saved": True}

@router.delete("/account/profile")
def delete_profile(user=Depends(current_user)):
    supabase.table("career_profiles").delete().eq("user_id", user).execute()
    return {"deleted": True}

@router.post("/account/cv/extract")
async def extract(file: UploadFile = File(...), user=Depends(current_user)):
    raw = await file.read(5 * 1024 * 1024 + 1)
    await file.close()
    if len(raw) > 5 * 1024 * 1024: raise HTTPException(413, "Upload a file under 5 MB.")
    extension = (file.filename or "").rsplit(".", 1)[-1].lower()
    try:
        if extension == "pdf":
            from pypdf import PdfReader
            pdf = PdfReader(io.BytesIO(raw))
            if len(pdf.pages) > 30: raise HTTPException(400, "Use a CV of 30 pages or fewer.")
            content = "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif extension == "docx":
            import zipfile
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024: raise HTTPException(400, "Document is too large to process.")
            from docx import Document
            document = Document(io.BytesIO(raw))
            content = "\n".join(p.text for p in document.paragraphs)
            content += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for t in document.tables for row in t.rows)
        elif extension == "txt": content = raw.decode("utf-8")
        else: raise HTTPException(400, "Use PDF, Word (.docx), or plain text.")
    except HTTPException: raise
    except Exception: raise HTTPException(400, "Could not read this document. Try Word or paste its text.")
    finally: del raw
    content = content.strip()[:30000]
    if len(content) < 30: raise HTTPException(400, "No readable CV text found. Paste the text from a scanned CV.")
    return {"text": content, "message": "The upload was processed in memory and was not saved. Review the extracted text before continuing."}

class CVRequest(BaseModel):
    verified_text: str = Field(min_length=30, max_length=30000)
    job_description: str = Field(default="", max_length=30000)
    job_title: str = Field(default="", max_length=250)
    company_name: str = Field(default="", max_length=250)
    confirmed: bool

SKILLS = ["python", "javascript", "typescript", "react", "sql", "excel", "power bi", "tableau", "aws", "azure", "django", "fastapi", "java", "c++", "accounting", "sales", "marketing", "customer service", "communication", "project management", "leadership", "data analysis", "nursing", "teaching", "research", "cybersecurity", "networking", "git", "docker", "kubernetes", "figma", "design", "finance", "procurement", "logistics", "administration"]

def found_skills(text):
    return {skill for skill in SKILLS if re.search(r"(?<!\w)" + re.escape(skill) + r"(?!\w)", text.lower())}

def prepare_cv(data):
    if not data.confirmed: raise HTTPException(400, "Review and confirm your CV facts first.")
    lines = [line.strip() for line in data.verified_text.splitlines() if line.strip()]
    required = found_skills(data.job_description)
    evidence = found_skills(data.verified_text)
    # Highlight exact source lines. No generated experience, skills or achievements.
    relevant = [line for line in lines if found_skills(line) & required]
    document = "\n".join(lines)
    if relevant:
        document = "RELEVANT EXPERIENCE AND SKILLS\n" + "\n".join(dict.fromkeys(relevant)) + "\n\nFULL CAREER RECORD\n" + document
    return {"document": document, "matched": sorted(required & evidence), "missing": sorted(required - evidence), "unassessed": "Review qualifications, work authorization, location and required years of experience manually. This keyword check is not a hiring probability.", "mode": "Evidence-based formatting and tailoring. Only your confirmed text is used."}

@router.post("/account/cv/prepare")
def prepare(data: CVRequest, user=Depends(current_user)):
    return prepare_cv(data)

@router.post("/account/cv/cover-letter")
def cover_letter(data: CVRequest, user=Depends(current_user)):
    result = prepare_cv(data)
    if not data.job_title: raise HTTPException(400, "Enter the job title first.")
    lines = [line.strip() for line in data.verified_text.splitlines() if line.strip()]
    evidence = [line for line in lines if found_skills(line) & set(result["matched"])][:3]
    supporting = "\n\n".join(evidence)
    document = f"Dear Hiring Team,\n\nI am applying for the {data.job_title} role" + (f" at {data.company_name}" if data.company_name else "") + ".\n\n"
    if supporting: document += "The following experience from my CV is relevant to the posting:\n\n" + supporting + "\n\n"
    document += "Thank you for considering my application. I would welcome the opportunity to discuss my experience.\n\nKind regards,\n"
    return {"document": document, "message": "Add your name and review the letter before using it. Supporting statements are copied from your confirmed CV."}

class Export(BaseModel):
    text: str = Field(min_length=1, max_length=40000)
    title: str = Field(default="CV", max_length=100)
    format: Literal["pdf", "docx"]

@router.post("/account/cv/export")
def export(data: Export, user=Depends(current_user)):
    output = io.BytesIO()
    if data.format == "docx":
        from docx import Document
        document = Document()
        document.core_properties.title = data.title
        for line in data.text.splitlines(): document.add_paragraph(line)
        document.save(output)
        content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from xml.sax.saxutils import escape
        from pathlib import Path
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        if 'ZagSans' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('ZagSans', str(Path(__file__).parent / 'assets' / 'DejaVuSans.ttf')))
        styles = getSampleStyleSheet()
        styles['Normal'].fontName = 'ZagSans'
        styles['Normal'].leading = 15
        story = []
        for line in data.text.splitlines():
            story.append(Paragraph(escape(line), styles["Normal"]) if line else Spacer(1, 8))
        SimpleDocTemplate(output, title=data.title).build(story)
        content_type = "application/pdf"
    return Response(output.getvalue(), media_type=content_type, headers={"Content-Disposition": f'attachment; filename="zagjobsearch-{data.format}.{data.format}"', "Cache-Control": "no-store"})

class OAuthExchange(BaseModel):
    access_token: str = Field(min_length=10, max_length=6000)

@router.post("/auth/google")
def google_exchange(data: OAuthExchange):
    try:
        verified = supabase.auth.get_user(data.access_token).user
        if not verified or not verified.email_confirmed_at: raise ValueError()
        provider = (verified.app_metadata or {}).get("provider")
        if provider != "google": raise ValueError()
        rows = supabase.table("profiles").select("id,email,full_name,country").eq("email", verified.email).execute().data
        if rows: profile = rows[0]
        else:
            raise HTTPException(409, "Create your account with this email first, then use Google to sign in.")
        return {**profile, "access_token": issue_session(profile["id"])}
    except HTTPException: raise
    except Exception: raise HTTPException(401, "Google sign-in could not be verified.")
